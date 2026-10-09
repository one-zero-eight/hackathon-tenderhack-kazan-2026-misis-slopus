#!/usr/bin/env python3
"""Read SKU IDs from given xlsx and download cards via ste_ui_fetch QueryIndex (batched)."""
from __future__ import annotations

import argparse
import logging
import sys
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ste_ui_fetch as ui

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
LOG_DIR = DATA / "logs"
SOURCES = ("Инструменты", "Мебель", "Одежда")


def setup_logging(stem: str | None = None) -> logging.Logger:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("download_ste_from_given_data")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    fmt = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    stream = logging.StreamHandler(sys.stdout)
    stream.setFormatter(fmt)
    logger.addHandler(stream)
    log_path = LOG_DIR / ("all.log" if stem is None else f"{stem}.log")
    file_handler = logging.FileHandler(log_path, encoding="utf-8")
    file_handler.setFormatter(fmt)
    logger.addHandler(file_handler)
    if stem is not None:
        combined = logging.FileHandler(LOG_DIR / "all.log", encoding="utf-8")
        combined.setFormatter(fmt)
        logger.addHandler(combined)
    return logger


def load_ids(xlsx: Path) -> list[int]:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        rows = wb[wb.sheetnames[0]].iter_rows(min_col=1, max_col=1, values_only=True)
        header = next(rows, None)
        if not header or str(header[0] or "").strip() != "ID СТЕ":
            raise RuntimeError(f"{xlsx.name}: expected header 'ID СТЕ', got {header!r}")
        ids, seen = [], set()
        for (value,) in rows:
            if value is None or str(value).strip() == "":
                continue
            sku_id = int(str(value).strip())
            if sku_id not in seen:
                seen.add(sku_id)
                ids.append(sku_id)
        return ids
    finally:
        wb.close()


def run_stem(stem: str, bridge: str | None, pause: float, workers: int, batch_size: int, limit: int | None) -> str:
    logger = setup_logging(stem)
    source = DATA / f"{stem}.xlsx"
    try:
        ids = load_ids(source)
        if limit is not None:
            ids = ids[:limit]
        logger.info("%s: %s IDs from %s workers=%s batch=%s", stem, len(ids), source.name, workers, batch_size)
        ui.resume_download_query_index(
            ids,
            DATA / f"{stem}-details.sqlite",
            DATA / f"{stem}_from_api.jsonl",
            bridge=bridge,
            pause=pause,
            workers=workers,
            batch_size=batch_size,
            list_source=str(source),
        )
        logger.info("%s: finished ok", stem)
        return f"{stem}:ok"
    except ui.BotProtected as e:
        logger.error("AGENT_ALERT kind=bot_protection stem=%s detail=%s", stem, e)
        raise
    except Exception as e:
        logger.exception("AGENT_ALERT kind=stem_crash stem=%s detail=%s", stem, e)
        raise


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only", choices=SOURCES, action="append")
    parser.add_argument("--bridge", default=None)
    parser.add_argument("--pause", type=float, default=0.0)
    parser.add_argument("--workers", type=int, default=4, help="Threads per process (use 4 when --parallel)")
    parser.add_argument("--batch-size", type=int, default=50)
    parser.add_argument("--limit", type=int, default=None, help="Download only first N IDs (smoke test)")
    parser.add_argument("--parallel", action="store_true", help="One process per xlsx")
    args = parser.parse_args()
    stems = tuple(args.only) if args.only else SOURCES
    setup_logging()
    log = logging.getLogger("download_ste_from_given_data")
    log.info("start stems=%s parallel=%s workers=%s batch=%s", stems, args.parallel, args.workers, args.batch_size)

    if args.parallel and len(stems) > 1:
        with ProcessPoolExecutor(max_workers=len(stems)) as pool:
            futs = {
                pool.submit(run_stem, stem, args.bridge, args.pause, args.workers, args.batch_size, args.limit): stem
                for stem in stems
            }
            for fut in as_completed(futs):
                stem = futs[fut]
                try:
                    log.info("%s", fut.result())
                except Exception as e:
                    log.error("AGENT_ALERT kind=process_fail stem=%s detail=%s", stem, e)
                    raise
    else:
        for stem in stems:
            run_stem(stem, args.bridge, args.pause, args.workers, args.batch_size, args.limit)

    log.info("all stems complete")


if __name__ == "__main__":
    main()
