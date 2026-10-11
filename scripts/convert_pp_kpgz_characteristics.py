#!/usr/bin/env python3
"""Parse Лист2 from portal STE export workbooks into pp-kpgz-characteristics.parquet.

Each row is one characteristic of a leaf КПГЗ / конечной категории ПП:
kind ∈ {Категореобразующая, Базовая, Эталонная}, typed values, optional СПГЗ link.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
DEFAULT_SOURCE = Path("/home/dante/Downloads/Казань 2026 (2)/Примеры выгрузок")
VALID_KINDS = frozenset({"Категореобразующая", "Базовая", "Эталонная"})
VALUE_SPLIT = re.compile(r"\s*;\s*")

SCHEMA = pa.schema([
    ("source_file", pa.string()),
    ("category_pp", pa.string()),
    ("kpgz_code", pa.string()),
    ("kpgz_name", pa.string()),
    ("name", pa.string()),
    ("kind", pa.string()),
    ("value_type", pa.string()),
    ("values", pa.list_(pa.string())),
    ("unit", pa.string()),
    ("spgz_characteristic_id", pa.int64()),
    ("spgz_characteristic_name", pa.string()),
    ("spgz_values", pa.list_(pa.string())),
])


def cell_str(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def cell_int(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    if text.isdigit():
        return int(text)
    return None


def split_values(value: object) -> list[str]:
    text = cell_str(value)
    if not text:
        return []
    parts = [part.strip() for part in VALUE_SPLIT.split(text) if part.strip()]
    return parts or [text]


def iter_sheet_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Лист2" not in workbook.sheetnames:
            raise ValueError(f"{path.name}: нет листа «Лист2»")
        sheet = workbook["Лист2"]
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index == 0:
                continue
            values = list(row)
            if not any(values):
                continue
            yield values
    finally:
        workbook.close()


def parse_row(source_file: str, row: list) -> dict | None:
    kind = cell_str(row[4] if len(row) > 4 else None)
    if kind not in VALID_KINDS:
        return None
    code = cell_str(row[1] if len(row) > 1 else None)
    name = cell_str(row[3] if len(row) > 3 else None)
    if not code or not name:
        return None
    return {
        "source_file": source_file,
        "category_pp": cell_str(row[0] if len(row) > 0 else None),
        "kpgz_code": code,
        "kpgz_name": cell_str(row[2] if len(row) > 2 else None),
        "name": name,
        "kind": kind,
        "value_type": cell_str(row[5] if len(row) > 5 else None),
        "values": split_values(row[6] if len(row) > 6 else None),
        "unit": cell_str(row[7] if len(row) > 7 else None),
        "spgz_characteristic_id": cell_int(row[8] if len(row) > 8 else None),
        "spgz_characteristic_name": cell_str(row[9] if len(row) > 9 else None),
        "spgz_values": split_values(row[10] if len(row) > 10 else None),
    }


def convert(sources: list[Path], output: Path) -> dict:
    records: list[dict] = []
    skipped = 0
    for path in sources:
        before = len(records)
        for row in iter_sheet_rows(path):
            parsed = parse_row(path.name, row)
            if parsed is None:
                skipped += 1
                continue
            records.append(parsed)
        print(f"{path.name}: {len(records) - before} rows")
    if not records:
        raise SystemExit("Нет валидных строк Лист2")
    # Stable order for reproducible parquet.
    records.sort(key=lambda item: (item["kpgz_code"], item["kind"], item["name"], item["source_file"]))
    table = pa.Table.from_pylist(records, schema=SCHEMA)
    output.parent.mkdir(parents=True, exist_ok=True)
    pq.write_table(table, output, compression="zstd")
    codes = {item["kpgz_code"] for item in records}
    return {"rows": len(records), "codes": len(codes), "skipped": skipped, "output": str(output)}


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source", type=Path, nargs="*", default=None,
        help="xlsx файлы или каталог с ними (по умолчанию примеры выгрузок Казани)",
    )
    parser.add_argument("--output", type=Path, default=DATA / "pp-kpgz-characteristics.parquet")
    args = parser.parse_args(argv)
    if not args.source:
        roots = [DEFAULT_SOURCE]
    else:
        roots = args.source
    files: list[Path] = []
    for root in roots:
        if root.is_dir():
            files.extend(sorted(root.glob("*.xlsx")))
        else:
            files.append(root)
    files = [path for path in files if path.is_file() and not path.name.startswith(".~")]
    if not files:
        raise SystemExit(f"Не найдены xlsx: {roots}")
    stats = convert(files, args.output)
    print(
        f"wrote {stats['rows']} rows / {stats['codes']} kpgz codes "
        f"(skipped {stats['skipped']}) → {stats['output']}"
    )


if __name__ == "__main__":
    main()
