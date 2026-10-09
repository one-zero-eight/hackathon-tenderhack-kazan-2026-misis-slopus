#!/usr/bin/env python3
"""Parse data/{Инструменты,Мебель,Одежда}.xlsx into <name>_from_data.jsonl."""
from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
SOURCES = ("Инструменты", "Мебель", "Одежда")

VALUE_START = re.compile(
    r"\s(?:\d+(?:\.\d+)?(?:-\d+(?:\.\d+)?)?|\((?:не задано)\))(?=\s|$)",
    re.IGNORECASE,
)
YES_NO = frozenset({"Да", "Нет", "да", "нет"})


def cell_str(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def candidate_key(part: str) -> str | None:
    m = VALUE_START.search(part)
    if m:
        key = part[: m.start()].strip()
        # Numbers inside comma-separated values are not key boundaries.
        if key and "," not in key:
            return key
    tokens = part.rsplit(None, 1)
    if len(tokens) == 2 and tokens[1] in YES_NO:
        return tokens[0].strip() or None
    return None


def value_starts_lower(value: str) -> bool:
    return bool(value) and value[0].isalpha() and value[0].islower()


def mine_string_keys(string_parts: list[str], min_suffixes: int = 3) -> set[str]:
    """Discover multi-word keys from free-text characteristics."""
    suffix_sets: dict[str, set[str]] = defaultdict(set)
    first_words: dict[str, set[str]] = defaultdict(set)
    for part in string_parts:
        words = part.split()
        for n in range(1, min(len(words), 8)):
            prefix = " ".join(words[:n])
            if "," in prefix:
                continue
            rest = words[n:]
            if not rest:
                continue
            suffix_sets[prefix].add(" ".join(rest))
            first_words[prefix].add(rest[0])
    keys: set[str] = set()
    for prefix, suffixes in suffix_sets.items():
        if len(suffixes) >= min_suffixes and len(first_words[prefix]) >= 2:
            keys.add(prefix)
    return keys


class KeyIndex:
    """Longest-prefix match of known characteristic names, indexed by first word."""

    def __init__(self, keys: set[str]):
        by_first: dict[str, list[str]] = {}
        for key in keys:
            first = key.split(None, 1)[0]
            by_first.setdefault(first, []).append(key)
        for group in by_first.values():
            group.sort(key=len, reverse=True)
        self._by_first = by_first

    def split(self, part: str) -> tuple[str, str]:
        first = part.split(None, 1)[0]
        matches: list[tuple[str, str]] = []
        for key in self._by_first.get(first, ()):
            if part == key:
                matches.append((key, ""))
            elif part.startswith(key) and len(part) > len(key) and part[len(key)].isspace():
                matches.append((key, part[len(key) :].strip()))
        if matches:
            with_value = [m for m in matches if m[1]]
            good = [m for m in with_value if not value_starts_lower(m[1])]
            if good:
                return max(good, key=lambda m: len(m[0]))
            if with_value:
                return max(with_value, key=lambda m: len(m[0]))
            return matches[0]

        key = candidate_key(part)
        if key is not None:
            return key, part[len(key) :].strip()
        tokens = part.split(None, 1)
        if len(tokens) == 2:
            return tokens[0], tokens[1]
        return part, ""


def collect_keys(parts: list[str]) -> set[str]:
    keys: set[str] = set()
    string_parts: list[str] = []
    for part in parts:
        key = candidate_key(part)
        if key:
            keys.add(key)
        else:
            string_parts.append(part)
    keys.update(mine_string_keys(string_parts))
    return keys


def parse_characteristics(raw: str | None, index: KeyIndex) -> list[dict[str, str]]:
    if not raw:
        return []
    out: list[dict[str, str]] = []
    for part in raw.split("|"):
        part = part.strip()
        if not part:
            continue
        key, value = index.split(part)
        out.append({"key": key, "value": value})
    return out


def iter_rows(xlsx: Path):
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            raise RuntimeError(f"{xlsx.name}: empty sheet")
        headers = [cell_str(h) for h in header_row]
        while headers and headers[-1] is None:
            headers.pop()
        if not headers or headers[0] != "ID СТЕ":
            raise RuntimeError(f"{xlsx.name}: expected header 'ID СТЕ', got {headers!r}")
        if "Характеристики" not in headers:
            raise RuntimeError(f"{xlsx.name}: missing 'Характеристики' column")
        yield headers
        for row in rows:
            values = list(row[: len(headers)])
            if len(values) < len(headers):
                values.extend([None] * (len(headers) - len(values)))
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            yield values
    finally:
        wb.close()


def convert_stem(stem: str) -> int:
    source = DATA / f"{stem}.xlsx"
    out_path = DATA / f"{stem}_from_data.jsonl"
    if not source.exists():
        raise FileNotFoundError(source)

    print(f"{stem}: reading {source.name}", flush=True)
    rows_iter = iter_rows(source)
    headers = next(rows_iter)
    char_idx = headers.index("Характеристики")

    cached: list[list] = []
    parts: list[str] = []
    for values in rows_iter:
        cached.append(values)
        raw = values[char_idx]
        if raw:
            parts.extend(p.strip() for p in str(raw).split("|") if p.strip())

    keys = collect_keys(parts)
    index = KeyIndex(keys)
    print(f"{stem}: {len(cached)} rows, {len(keys)} characteristic keys", flush=True)

    count = 0
    with out_path.open("w", encoding="utf-8") as out:
        for values in cached:
            record = {}
            for header, value in zip(headers, values):
                if header is None:
                    continue
                if header == "Характеристики":
                    record[header] = parse_characteristics(cell_str(value), index)
                elif header == "ID СТЕ":
                    text = cell_str(value)
                    record[header] = int(text) if text is not None else None
                else:
                    record[header] = cell_str(value)
            out.write(json.dumps(record, ensure_ascii=False) + "\n")
            count += 1
            if count % 20000 == 0:
                print(f"{stem}: {count} written...", flush=True)

    print(f"{stem}: wrote {count} -> {out_path.name}", flush=True)
    return count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only", choices=SOURCES, action="append")
    args = parser.parse_args()
    stems = tuple(args.only) if args.only else SOURCES
    total = 0
    for stem in stems:
        total += convert_stem(stem)
    print(f"done: {total} records", flush=True)


if __name__ == "__main__":
    main()
